"""WhatsApp template campaign, follow-up, and outbox service."""
from __future__ import annotations

import asyncio
import csv
import io
import logging
import uuid
from datetime import datetime, timedelta, timezone
from typing import Any

import httpx
from fastapi import UploadFile
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import settings
from app.core.exceptions import ExternalServiceError
from app.integrations.llm.client import get_llm_client
from app.integrations.meta_whatsapp.client import MetaWhatsAppClient
from app.models.company_config import CompanyConfig
from app.models.conversation import Conversation
from app.models.message import Message
from app.models.whatsapp_campaign import (
    WhatsAppCampaign,
    WhatsAppCampaignRecipient,
    WhatsAppFollowupRule,
    WhatsAppOutboxJob,
    WhatsAppSuppression,
)
from app.repositories.company_config_repository import CompanyConfigRepository
from app.repositories.conversation_repository import ConversationRepository
from app.repositories.message_repository import MessageRepository
from app.repositories.whatsapp_campaign_repository import (
    WhatsAppCampaignRepository,
    WhatsAppFollowupRuleRepository,
    WhatsAppOutboxRepository,
    WhatsAppRecipientRepository,
    WhatsAppSuppressionRepository,
)
from app.schemas.whatsapp_campaign import (
    CampaignCreateRequest,
    CampaignPreviewResponse,
    FollowupRuleCreateRequest,
    FollowupRuleUpdateRequest,
    MetaTemplateCreateRequest,
    QuickSendRequest,
)
from app.services.whatsapp_outbound import company_can_send_whatsapp, send_company_whatsapp_text
from app.utils.meta_whatsapp_template import build_simple_template_create_payload

logger = logging.getLogger(__name__)


_FOLLOWUP_AI_SYSTEM_PROMPT = (
    "You are drafting ONE WhatsApp outbound message — a polite, short follow-up to a customer whose "
    "support inquiry appears open.\n\n"
    "Rules:\n"
    "- Write ONLY the message body. No greetings that quote internal policies. No preamble.\n"
    "- Warm, respectful, WhatsApp-toned; fewer than ~600 characters when possible unless the conversation truly needs more.\n"
    "- Acknowledge ongoing help and gently ask if they still need assistance or had further questions.\n"
    "- Do not invent specifics (order numbers, prices, diagnoses, promises); stay generic unless the thread clearly states them.\n"
    "- Avoid medical diagnoses, legal certainty, refunds-as-guarantees, or pressuring urgency.\n"
    "- Respect any opt-out cues in history; if they declined contact, apologize briefly without selling.\n"
    "- Match conversation language strictly as instructed by LANGUAGE RULE snippets from the assistant layer."
)


def _detected_language_to_llm(detected: str | None) -> str:
    s = (detected or "english").strip().lower().replace("-", "_")
    if not s:
        return "english"
    if "hinglish" in s:
        return "hinglish"
    if s.startswith("hi") or s == "hindi":
        return "hindi"
    return "english"


OPT_OUT_KEYWORDS = {
    "stop",
    "unsubscribe",
    "unsub",
    "opt out",
    "opt-out",
    "cancel",
    "band karo",
    "band kardo",
    "mat bhejo",
    "message mat bhejo",
    # Extended English
    "please stop",
    "stop messaging",
    "stop sending",
    "no more messages",
    "remove me",
    "do not contact",
    "don't contact",
    "do not message",
    "don't message",
    "leave me alone",
    "dnd",
    "do not disturb",
    # Extended Hindi / Hinglish
    "message mat karo",
    "send mat karo",
    "mujhe mat bhejo",
    "rok o",
    "roko",
    "hatao",
    "nahi chahiye",
    "pareshan mat karo",
    "disturb mat karo",
}


def is_opt_out_text(text: str) -> bool:
    from app.utils.opt_out import is_opt_out_message
    raw = " ".join((text or "").strip().lower().split())
    if not raw:
        return False
    if raw in OPT_OUT_KEYWORDS or any(raw.startswith(k + " ") for k in OPT_OUT_KEYWORDS):
        return True
    # Fallback: regex-based detection for more complex phrases
    return is_opt_out_message(text)


def _meta_client(config: CompanyConfig) -> MetaWhatsAppClient:
    token = config.meta_graph_access_token
    phone_id = (config.meta_phone_number_id or "").strip()
    if not token or not phone_id:
        raise RuntimeError("Meta credentials are required for campaign sends.")
    return MetaWhatsAppClient(
        access_token=token,
        phone_number_id=phone_id,
        graph_base_url=settings.meta_graph_api_base_url,
        graph_version=settings.meta_graph_api_version,
        timeout=settings.meta_webhook_timeout_seconds,
    )


class WhatsAppCampaignService:
    def __init__(self, db: AsyncSession) -> None:
        self.db = db
        self.campaigns = WhatsAppCampaignRepository(db)
        self.recipients = WhatsAppRecipientRepository(db)
        self.outbox = WhatsAppOutboxRepository(db)
        self.followups = WhatsAppFollowupRuleRepository(db)
        self.suppressions = WhatsAppSuppressionRepository(db)

    async def preview_upload(self, file: UploadFile) -> CampaignPreviewResponse:
        content = await file.read()
        rows = _parse_tabular_file(file.filename or "", content)
        columns = list(rows[0].keys()) if rows else []
        errors: list[str] = []
        valid = 0
        phone_candidates = [c for c in columns if c.strip().lower() in {"phone", "phone_number", "mobile", "whatsapp", "number"}]
        phone_col = phone_candidates[0] if phone_candidates else (columns[0] if columns else "")
        preview_rows: list[dict[str, Any]] = []
        for idx, row in enumerate(rows, start=2):
            enriched = dict(row)
            raw_phone = str(row.get(phone_col, "") or "")
            try:
                enriched["_normalized_phone"] = normalize_whatsapp_destination(raw_phone)
                valid += 1
            except Exception as exc:
                enriched["_error"] = str(exc)
                errors.append(f"Row {idx}: {exc}")
            preview_rows.append(enriched)
        invalid = max(0, len(rows) - valid)
        return CampaignPreviewResponse(
            columns=columns,
            rows=preview_rows,
            total_rows=len(rows),
            valid_rows=valid,
            invalid_rows=invalid,
            errors=errors[:50],
        )

    async def create_campaign(self, company_id: uuid.UUID, body: CampaignCreateRequest) -> WhatsAppCampaign:
        campaign = await self.campaigns.create(
            {
                "company_id": company_id,
                "name": body.name.strip(),
                "status": "draft",
                "template_name": body.template_name.strip(),
                "language_code": body.language_code.strip(),
                "body_variable_mappings": body.body_variable_mappings,
                "header_media_url_mapping": body.header_media_url_mapping,
                "total_recipients": 0,
                "queued_count": 0,
                "sent_count": 0,
                "failed_count": 0,
                "skipped_count": 0,
            }
        )
        total = 0
        skipped = 0
        for idx, row in enumerate(body.rows, start=2):
            raw_phone = str(row.get(body.phone_column, "") or "")
            try:
                phone = normalize_whatsapp_destination(raw_phone)
                status = "queued"
                err = None
            except Exception as exc:
                phone = ""
                status = "skipped"
                err = str(exc)
                skipped += 1
            await self.recipients.create(
                {
                    "company_id": company_id,
                    "campaign_id": campaign.id,
                    "row_index": idx,
                    "phone_number": phone,
                    "raw_phone": raw_phone,
                    "row_json": row,
                    "status": status,
                    "error_message": err,
                }
            )
            total += 1
        await self.campaigns.update(campaign, {"total_recipients": total, "skipped_count": skipped})
        return campaign

    async def start_campaign(self, company_id: uuid.UUID, campaign_id: uuid.UUID) -> WhatsAppCampaign:
        campaign = await self._get_campaign(company_id, campaign_id)
        if campaign.status == "running":
            return campaign
        if campaign.status == "paused":
            return await self.campaigns.update(campaign, {"status": "running"})
        if campaign.status in {"completed", "cancelled", "template_rejected"}:
            raise ValueError(f"Campaign is {campaign.status} and cannot be started again.")

        recipients = await self.recipients.list_for_campaign(campaign.id, limit=100_000)
        queued = 0
        suppressed = 0
        for recipient in recipients:
            if recipient.status != "queued" or not recipient.phone_number:
                continue
            if recipient.outbox_job_id:
                queued += 1
                continue
            if await self.is_suppressed(company_id, recipient.phone_number):
                await self.recipients.update(recipient, {"status": "suppressed", "error_message": "Suppressed/opted out"})
                suppressed += 1
                continue
            variables = self._resolve_campaign_variables(recipient.row_json, campaign.body_variable_mappings or [])
            header_url = None
            if campaign.header_media_url_mapping:
                header_url = str(recipient.row_json.get(campaign.header_media_url_mapping, "") or "").strip() or None
            job = await self.outbox.create(
                {
                    "company_id": company_id,
                    "kind": "campaign",
                    "status": "queued",
                    "to_number": recipient.phone_number,
                    "template_name": campaign.template_name,
                    "language_code": campaign.language_code,
                    "body_variables_json": variables,
                    "header_media_url": header_url,
                    "payload_json": {"row": recipient.row_json},
                    "campaign_id": campaign.id,
                    "recipient_id": recipient.id,
                    "max_attempts": settings.whatsapp_outbox_retry_max_attempts,
                    "next_attempt_at": datetime.now(timezone.utc),
                }
            )
            await self.recipients.update(recipient, {"outbox_job_id": job.id})
            queued += 1
        await self.campaigns.update(
            campaign,
            {
                "status": "running",
                "queued_count": queued,
                "skipped_count": campaign.skipped_count + suppressed,
                "started_at": campaign.started_at or datetime.now(timezone.utc),
            },
        )
        return campaign

    async def campaign_action(self, company_id: uuid.UUID, campaign_id: uuid.UUID, action: str) -> WhatsAppCampaign:
        if action == "start":
            return await self.start_campaign(company_id, campaign_id)
        campaign = await self._get_campaign(company_id, campaign_id)
        if action == "pause":
            return await self.campaigns.update(campaign, {"status": "paused"})
        if action == "cancel":
            jobs = await self.outbox.list_for_company(company_id, campaign_id=campaign.id, limit=100_000)
            for job in jobs:
                if job.status == "queued":
                    job.status = "cancelled"
            return await self.campaigns.update(campaign, {"status": "cancelled", "completed_at": datetime.now(timezone.utc)})
        if action == "retry_failed":
            jobs = await self.outbox.list_for_company(company_id, campaign_id=campaign.id, status="failed", limit=100_000)
            for job in jobs:
                job.status = "queued"
                job.next_attempt_at = datetime.now(timezone.utc)
                job.last_error = None
            await self.db.flush()
            return campaign
        raise ValueError(f"Unknown action {action}")

    async def list_campaigns(self, company_id: uuid.UUID) -> list[WhatsAppCampaign]:
        return await self.campaigns.list_for_company(company_id)

    async def campaign_detail(self, company_id: uuid.UUID, campaign_id: uuid.UUID) -> tuple[WhatsAppCampaign, list[WhatsAppCampaignRecipient], dict[str, int]]:
        campaign = await self._get_campaign(company_id, campaign_id)
        return campaign, await self.recipients.list_for_campaign(campaign_id), await self.outbox.counts_for_campaign(campaign_id)

    async def list_jobs(self, company_id: uuid.UUID, campaign_id: uuid.UUID | None = None, status: str | None = None) -> list[WhatsAppOutboxJob]:
        return await self.outbox.list_for_company(company_id, campaign_id=campaign_id, status=status)

    async def create_followup_rule(self, company_id: uuid.UUID, body: FollowupRuleCreateRequest) -> WhatsAppFollowupRule:
        return await self.followups.create(
            {"company_id": company_id, "name": body.name.strip(), "is_active": body.is_active, "steps_json": [s.model_dump() for s in body.steps]}
        )

    async def update_followup_rule(self, company_id: uuid.UUID, rule_id: uuid.UUID, body: FollowupRuleUpdateRequest) -> WhatsAppFollowupRule:
        rule = await self.followups.get(rule_id)
        if not rule or rule.company_id != company_id:
            raise ValueError("Follow-up rule not found.")
        updates: dict[str, Any] = {}
        if body.name is not None:
            updates["name"] = body.name.strip()
        if body.is_active is not None:
            updates["is_active"] = body.is_active
        if body.steps is not None:
            updates["steps_json"] = [s.model_dump() for s in body.steps]
        return await self.followups.update(rule, updates)

    async def delete_followup_rule(self, company_id: uuid.UUID, rule_id: uuid.UUID) -> None:
        rule = await self.followups.get(rule_id)
        if not rule or rule.company_id != company_id:
            raise ValueError("Follow-up rule not found.")
        await self.followups.delete(rule)

    async def schedule_due_followups(self, company_id: uuid.UUID | None = None) -> int:
        rules = await self.followups.list_active(company_id)
        created = 0
        for rule in rules:
            result = await self.db.execute(
                select(Conversation).where(
                    Conversation.company_id == rule.company_id,
                    Conversation.inquiry_complete.is_(False),
                    Conversation.is_blocked.is_(False),
                    Conversation.opted_out.is_(False),
                )
            )
            conversations = list(result.scalars().all())
            for conv in conversations:
                if await self.is_suppressed(conv.company_id, conv.customer_phone):
                    continue
                if await self._customer_replied_after_last_followup(conv.id, rule.id):
                    continue
                last_customer = await self._last_customer_message_at(conv.id)
                if last_customer is None:
                    continue
                for idx, step in enumerate((rule.steps_json or [])[:3]):
                    if await self.outbox.exists_followup_job(
                        conversation_id=conv.id, followup_rule_id=rule.id, step_index=idx
                    ):
                        continue
                    due = last_customer + timedelta(minutes=int(step.get("delay_minutes") or 0))
                    if due > datetime.now(timezone.utc):
                        continue
                    await self.outbox.create(
                        {
                            "company_id": conv.company_id,
                            "kind": "followup",
                            "status": "queued",
                            "to_number": conv.customer_phone,
                            "template_name": str(step.get("template_name") or ""),
                            "language_code": str(step.get("language_code") or "en"),
                            "body_variables_json": [str(v) for v in step.get("body_variables") or []],
                            "header_media_url": step.get("header_media_url"),
                            "payload_json": {"conversation_id": str(conv.id)},
                            "followup_rule_id": rule.id,
                            "conversation_id": conv.id,
                            "followup_step_index": idx,
                            "max_attempts": settings.whatsapp_outbox_retry_max_attempts,
                            "next_attempt_at": datetime.now(timezone.utc),
                        }
                    )
                    created += 1
        return created

    async def process_due_outbox_once(self) -> int:
        await self.schedule_due_followups()
        jobs = await self.outbox.list_due(settings.whatsapp_outbox_batch_size)
        processed = 0
        last_company: uuid.UUID | None = None
        for job in jobs:
            if last_company == job.company_id:
                await asyncio.sleep(max(0, settings.whatsapp_outbox_company_delay_seconds))
            await self._process_job(job)
            last_company = job.company_id
            processed += 1
        return processed

    async def suppress_phone(self, company_id: uuid.UUID, phone_number: str, *, reason: str = "opt_out", source: str = "whatsapp") -> WhatsAppSuppression:
        existing = await self.suppressions.get_for_phone(company_id, phone_number)
        if existing:
            return existing
        return await self.suppressions.create(
            {"company_id": company_id, "phone_number": phone_number, "reason": reason, "source": source}
        )

    async def is_suppressed(self, company_id: uuid.UUID, phone_number: str) -> bool:
        return await self.suppressions.get_for_phone(company_id, phone_number) is not None

    def _resolve_campaign_variables(self, row: dict[str, Any], mappings: list[Any]) -> list[str]:
        variables: list[str] = []
        for mapping in mappings:
            if isinstance(mapping, dict):
                source = str(mapping.get("source") or "column").strip().lower()
                if source == "literal":
                    variables.append(str(mapping.get("value") or ""))
                    continue
                column = str(mapping.get("column") or mapping.get("name") or "").strip()
                variables.append(str(row.get(column, "") or "") if column else "")
                continue
            column = str(mapping or "").strip()
            variables.append(str(row.get(column, "") or "") if column else "")
        return variables

    async def _resolve_meta_waba_id(self, config: CompanyConfig) -> str | None:
        raw = getattr(config, "meta_waba_id", None)
        if raw is not None:
            s = str(raw).strip()
            if s:
                return s
        try:
            client = _meta_client(config)
        except RuntimeError:
            return None
        try:
            return await client.fetch_whatsapp_business_account_id()
        except ExternalServiceError:
            logger.warning(
                "Could not resolve Meta WABA id from phone_number_id",
                extra={"company_id": str(config.company_id)},
            )
            return None

    async def list_templates(self, company_id: uuid.UUID) -> list[dict[str, Any]]:
        config = await CompanyConfigRepository(self.db).get_by_company(company_id)
        if not config:
            return []
        waba_id = await self._resolve_meta_waba_id(config)
        if not waba_id:
            return []
        return await _meta_client(config).list_message_templates(waba_id=waba_id)

    async def create_meta_template(
        self,
        company_id: uuid.UUID,
        body: MetaTemplateCreateRequest,
    ) -> dict[str, Any]:
        config = await CompanyConfigRepository(self.db).get_by_company(company_id)
        if not config:
            raise ValueError("Company configuration missing.")
        waba_id = await self._resolve_meta_waba_id(config)
        if not waba_id:
            raise ValueError(
                "Could not resolve WhatsApp Business Account id. "
                "Save Meta phone_number_id and Graph token (Business Management permission)."
            )
        try:
            client = _meta_client(config)
        except RuntimeError as exc:
            raise ValueError(str(exc)) from exc

        payload = build_simple_template_create_payload(
            name=body.name,
            category=body.category,
            language=body.language,
            body_text=body.body_text,
            footer_text=body.footer_text,
            header_text=body.header_text,
            body_example_values=list(body.body_example_values),
        )
        return await client.create_message_template(waba_id=waba_id, payload=payload)

    async def create_campaign_with_template(
        self,
        company_id: uuid.UUID,
        body: "QuickSendRequest",
    ) -> tuple["WhatsAppCampaign", str]:
        """Submit a Meta template and create a campaign in one call.

        Returns ``(campaign, template_status)`` where ``template_status`` is the
        Meta status string (``APPROVED``, ``PENDING``, ``IN_APPEAL`` …).

        If the template is already APPROVED the campaign is immediately started.
        Otherwise it is left in ``pending_template`` status and the outbox worker
        will auto-start it once Meta approves.
        """
        config = await CompanyConfigRepository(self.db).get_by_company(company_id)
        if not config:
            raise ValueError("Company configuration missing.")
        waba_id = await self._resolve_meta_waba_id(config)
        if not waba_id:
            raise ValueError("Could not resolve WhatsApp Business Account ID.")

        client = _meta_client(config)

        # Check if template already exists and is approved
        existing_templates = await client.list_message_templates(waba_id=waba_id)
        existing = next(
            (t for t in existing_templates if t.get("name") == body.template_name),
            None,
        )
        if existing and existing.get("status") == "APPROVED":
            template_status = "APPROVED"
        else:
            # Submit/re-submit to Meta
            payload = build_simple_template_create_payload(
                name=body.template_name,
                category=body.category,
                language=body.language_code,
                body_text=body.template_body,
                footer_text=body.template_footer,
                header_text=body.template_header,
                body_example_values=[],
            )
            result = await client.create_message_template(waba_id=waba_id, payload=payload)
            template_status = result.get("status", "PENDING")

        initial_status = "draft" if template_status == "APPROVED" else "pending_template"

        create_body = CampaignCreateRequest(
            name=body.campaign_name,
            template_name=body.template_name,
            language_code=body.language_code,
            phone_column=body.phone_column,
            rows=body.rows,
        )
        campaign = await self.create_campaign(company_id, create_body)
        campaign = await self.campaigns.update(campaign, {"status": initial_status})

        if template_status == "APPROVED":
            campaign = await self.start_campaign(company_id, campaign.id)

        return campaign, template_status

    async def auto_start_pending_campaigns(self) -> int:
        """Check Meta template approval for all pending_template campaigns and start approved ones.

        Called periodically by the outbox worker. Returns number of campaigns started.
        """
        result = await self.db.execute(
            select(WhatsAppCampaign).where(WhatsAppCampaign.status == "pending_template")
        )
        pending = list(result.scalars().all())
        if not pending:
            return 0

        started = 0
        for campaign in pending:
            try:
                config = await CompanyConfigRepository(self.db).get_by_company(campaign.company_id)
                if not config:
                    continue
                waba_id = await self._resolve_meta_waba_id(config)
                if not waba_id:
                    continue
                client = _meta_client(config)
                templates = await client.list_message_templates(waba_id=waba_id)
                match = next(
                    (t for t in templates if t.get("name") == campaign.template_name),
                    None,
                )
                if not match:
                    continue
                status = (match.get("status") or "").upper()
                if status == "APPROVED":
                    await self.start_campaign(campaign.company_id, campaign.id)
                    started += 1
                    logger.info(
                        "Auto-started campaign after template approval",
                        extra={"campaign_id": str(campaign.id), "template": campaign.template_name},
                    )
                elif status == "REJECTED":
                    await self.campaigns.update(campaign, {"status": "template_rejected"})
                    logger.warning(
                        "Campaign template rejected by Meta",
                        extra={"campaign_id": str(campaign.id), "template": campaign.template_name},
                    )
            except Exception as exc:
                logger.warning(
                    "auto_start_pending_campaigns: error checking campaign",
                    extra={"campaign_id": str(campaign.id), "error": str(exc)},
                )
        return started

    async def _generate_followup_message(self, conversation: Conversation) -> str:
        """One LLM-produced follow-up grounded on recent transcript (no RAG)."""
        msg_repo = MessageRepository(self.db)
        rows = await msg_repo.list_recent_for_conversation(conversation.id, limit=28)
        lines: list[str] = []
        for m in rows:
            txt = ((m.message_text or "").strip())[:900]
            if not txt:
                continue
            st = (m.sender_type or "unknown").strip().upper()
            lines.append(f"{st}: {txt}")
        history_raw = "\n".join(lines)
        history_clip = history_raw[:3800]

        stub_ctx = "(No retrieval context — base the follow-up ONLY on conversation history.)"
        lang = _detected_language_to_llm(getattr(conversation, "detected_language", None))
        llm = get_llm_client()
        return (
            await llm.generate_answer(
                [stub_ctx],
                user_query=(
                    "Write the ONE WhatsApp follow-up message — output ONLY the customer's message "
                    "(no prefixes like 'Assistant:')."
                ),
                system_prompt=_FOLLOWUP_AI_SYSTEM_PROMPT,
                output_language=lang,
                conversation_history=history_clip if history_clip.strip() else None,
            )
        ).strip()

    async def _process_job(self, job: WhatsAppOutboxJob) -> None:
        config = await CompanyConfigRepository(self.db).get_by_company(job.company_id)
        if not config:
            await self._fail_or_retry(
                job,
                RuntimeError("Company configuration missing."),
                retryable=False,
            )
            return

        meta_ready = bool(config.meta_graph_access_token and (config.meta_phone_number_id or "").strip())

        if await self.is_suppressed(job.company_id, job.to_number):
            await self.outbox.update(job, {"status": "suppressed", "last_error": "Suppressed/opted out"})
            await self._sync_recipient_from_job(job, "suppressed", "Suppressed/opted out")
            await self._refresh_campaign_counts(job.campaign_id)
            return

        conv: Conversation | None = None

        if job.kind == "campaign":
            if not meta_ready:
                await self._fail_or_retry(
                    job,
                    RuntimeError("Meta WhatsApp credentials missing for bulk template campaign."),
                    retryable=False,
                )
                return
            if job.campaign_id:
                campaign = await self.campaigns.get(job.campaign_id)
                if campaign and campaign.status == "paused":
                    await self.outbox.update(
                        job,
                        {
                            "status": "queued",
                            "next_attempt_at": datetime.now(timezone.utc) + timedelta(seconds=60),
                            "last_error": "Campaign paused",
                        },
                    )
                    return
                if campaign and campaign.status == "cancelled":
                    await self.outbox.update(job, {"status": "skipped", "last_error": "Campaign cancelled"})
                    await self._sync_recipient_from_job(job, "skipped", f"Campaign {campaign.status}")
                    await self._refresh_campaign_counts(job.campaign_id)
                    return

        if job.kind == "followup" and job.conversation_id:
            if job.followup_rule_id:
                rule = await self.followups.get(job.followup_rule_id)
                if not rule or not rule.is_active:
                    await self.outbox.update(job, {"status": "skipped", "last_error": "Follow-up rule inactive or missing"})
                    return
            conv = await self.db.get(Conversation, job.conversation_id)
            if not conv or conv.inquiry_complete:
                await self.outbox.update(job, {"status": "skipped", "last_error": "Conversation complete or missing"})
                return
            if job.followup_rule_id and await self._customer_replied_after_last_followup(conv.id, job.followup_rule_id):
                await self.outbox.update(job, {"status": "skipped", "last_error": "Customer replied after latest follow-up"})
                return

        followup_ai_body: str | None = None
        if job.kind == "followup" and conv is not None:
            last_customer = await self._last_customer_message_at(conv.id)
            window = timedelta(hours=float(settings.whatsapp_followup_session_hours))
            in_session = (
                last_customer is not None
                and datetime.now(timezone.utc) <= last_customer + window
            )
            if in_session and company_can_send_whatsapp(config):
                try:
                    generated = await self._generate_followup_message(conv)
                except Exception:
                    logger.exception(
                        "Follow-up AI compose failed",
                        extra={"conversation_id": str(conv.id), "job_id": str(job.id)},
                    )
                    generated = ""
                clipped = (generated or "").strip()
                if clipped:
                    followup_ai_body = clipped[:4000]

        await self.outbox.update(job, {"status": "sending", "attempts": job.attempts + 1})
        try:
            if followup_ai_body:
                if conv is None:
                    raise RuntimeError("Follow-up AI path requires a loaded conversation.")
                await send_company_whatsapp_text(
                    config=config,
                    to_number=job.to_number,
                    text=followup_ai_body,
                )
                msg_repo = MessageRepository(self.db)
                conv_repo = ConversationRepository(self.db)
                now = datetime.now(timezone.utc)
                await msg_repo.create(
                    {
                        "conversation_id": conv.id,
                        "company_id": conv.company_id,
                        "sender_type": "bot",
                        "message_text": followup_ai_body,
                        "normalized_text": None,
                        "language": getattr(conv, "detected_language", None),
                        "response_type": None,
                        "external_message_id": None,
                    }
                )
                await conv_repo.touch_last_message_at(conv.id, now)
                await self.outbox.update(
                    job,
                    {
                        "status": "sent",
                        "sent_at": now,
                        "provider_message_id": None,
                        "last_error": None,
                    },
                )
                await self._refresh_campaign_counts(job.campaign_id)
                return

            if not meta_ready:
                await self._fail_or_retry(
                    job,
                    RuntimeError(
                        "Meta WhatsApp credentials required for template sends "
                        "(or enable a session-capable outbound provider)."
                    ),
                    retryable=False,
                )
                return

            data = await _meta_client(config).send_template(
                to_number=job.to_number,
                template_name=job.template_name,
                language_code=job.language_code,
                body_variables=[str(v) for v in (job.body_variables_json or [])],
                header_media_url=job.header_media_url,
            )
            msg_id = _provider_message_id(data)
            await self.outbox.update(
                job,
                {"status": "sent", "sent_at": datetime.now(timezone.utc), "provider_message_id": msg_id, "last_error": None},
            )
            await self._sync_recipient_from_job(job, "sent", None)
            await self._refresh_campaign_counts(job.campaign_id)
        except httpx.HTTPError as exc:
            await self._fail_or_retry(job, exc, retryable=_is_retryable(exc))
        except ExternalServiceError as exc:
            await self._fail_or_retry(job, exc, retryable=_is_retryable_external(exc))
        except Exception as exc:
            await self._fail_or_retry(job, exc, retryable=True)

    async def _fail_or_retry(self, job: WhatsAppOutboxJob, exc: Exception, *, retryable: bool) -> None:
        if retryable and job.attempts < job.max_attempts:
            delay = settings.whatsapp_outbox_retry_base_seconds * (2 ** max(job.attempts - 1, 0))
            await self.outbox.update(
                job,
                {"status": "queued", "next_attempt_at": datetime.now(timezone.utc) + timedelta(seconds=delay), "last_error": str(exc)[:2000]},
            )
        else:
            await self.outbox.update(job, {"status": "failed", "last_error": str(exc)[:2000]})
            await self._sync_recipient_from_job(job, "failed", str(exc)[:1000])
            await self._refresh_campaign_counts(job.campaign_id)

    async def _sync_recipient_from_job(self, job: WhatsAppOutboxJob, status: str, error: str | None) -> None:
        if not job.recipient_id:
            return
        recipient = await self.recipients.get(job.recipient_id)
        if recipient:
            await self.recipients.update(recipient, {"status": status, "error_message": error})

    async def _refresh_campaign_counts(self, campaign_id: uuid.UUID | None) -> None:
        if not campaign_id:
            return
        campaign = await self.campaigns.get(campaign_id)
        if not campaign:
            return
        counts = await self.outbox.counts_for_campaign(campaign_id)
        sent = counts.get("sent", 0)
        failed = counts.get("failed", 0)
        skipped = counts.get("skipped", 0) + counts.get("suppressed", 0) + counts.get("cancelled", 0)
        queued = counts.get("queued", 0) + counts.get("sending", 0)
        status_update: dict[str, Any] = {
            "sent_count": sent,
            "failed_count": failed,
            "skipped_count": skipped,
            "queued_count": queued,
        }
        if campaign.status == "running" and queued == 0:
            status_update["status"] = "completed"
            status_update["completed_at"] = datetime.now(timezone.utc)
        await self.campaigns.update(campaign, status_update)

    async def _get_campaign(self, company_id: uuid.UUID, campaign_id: uuid.UUID) -> WhatsAppCampaign:
        campaign = await self.campaigns.get(campaign_id)
        if not campaign or campaign.company_id != company_id:
            raise ValueError("Campaign not found.")
        return campaign

    async def _last_customer_message_at(self, conversation_id: uuid.UUID) -> datetime | None:
        result = await self.db.execute(
            select(Message.created_at)
            .where(Message.conversation_id == conversation_id, Message.sender_type == "customer")
            .order_by(Message.created_at.desc())
            .limit(1)
        )
        return result.scalar_one_or_none()

    async def _customer_replied_after_last_followup(
        self,
        conversation_id: uuid.UUID,
        followup_rule_id: uuid.UUID,
    ) -> bool:
        last_customer = await self._last_customer_message_at(conversation_id)
        if last_customer is None:
            return False
        result = await self.db.execute(
            select(WhatsAppOutboxJob.sent_at)
            .where(
                WhatsAppOutboxJob.conversation_id == conversation_id,
                WhatsAppOutboxJob.followup_rule_id == followup_rule_id,
                WhatsAppOutboxJob.kind == "followup",
                WhatsAppOutboxJob.status == "sent",
                WhatsAppOutboxJob.sent_at.is_not(None),
            )
            .order_by(WhatsAppOutboxJob.sent_at.desc())
            .limit(1)
        )
        latest_followup = result.scalar_one_or_none()
        return bool(latest_followup and last_customer > latest_followup)


def _parse_tabular_file(filename: str, content: bytes) -> list[dict[str, Any]]:
    lower = filename.lower()
    if lower.endswith(".csv"):
        text = content.decode("utf-8-sig")
        return [dict(row) for row in csv.DictReader(io.StringIO(text))]
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    headers = [str(v).strip() if v is not None else "" for v in next(rows_iter, [])]
    out: list[dict[str, Any]] = []
    for row in rows_iter:
        item = {headers[i]: row[i] for i in range(min(len(headers), len(row))) if headers[i]}
        if any(v is not None and str(v).strip() for v in item.values()):
            out.append(item)
    return out


def _is_retryable(exc: httpx.HTTPError) -> bool:
    if isinstance(exc, httpx.HTTPStatusError):
        return exc.response.status_code in {429, 500, 502, 503, 504}
    return isinstance(exc, httpx.RequestError)


def _is_retryable_external(exc: ExternalServiceError) -> bool:
    text = str(exc)
    return any(f"HTTP {code}" in text for code in (429, 500, 502, 503, 504))


def _provider_message_id(data: dict[str, Any]) -> str | None:
    messages = data.get("messages")
    if isinstance(messages, list) and messages:
        first = messages[0]
        if isinstance(first, dict):
            return str(first.get("id") or "") or None
    return None
